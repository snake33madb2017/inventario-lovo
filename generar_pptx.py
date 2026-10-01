import collections 
import collections.abc
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os

prs = Presentation()

# Colores y estilo LIMPIO / CLARO
BG_COLOR = RGBColor(250, 250, 250)
TEXT_COLOR = RGBColor(50, 50, 50)
ACCENT_COLOR = RGBColor(41, 128, 185) # Azul profesional

def add_slide(prs, title_text, content_text, image_path=None):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    
    # Fondo
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR
    
    # Título
    title_shape = slide.shapes.title
    title_shape.text = title_text
    title_shape.text_frame.paragraphs[0].font.color.rgb = ACCENT_COLOR
    title_shape.text_frame.paragraphs[0].font.bold = True
    title_shape.text_frame.paragraphs[0].font.name = "Arial"
    
    # Contenido
    body_shape = slide.placeholders[1]
    tf = body_shape.text_frame
    tf.text = content_text
    for paragraph in tf.paragraphs:
        paragraph.font.color.rgb = TEXT_COLOR
        paragraph.font.size = Pt(18)
        paragraph.font.name = "Arial"
        
    # Imagen si existe
    if image_path and os.path.exists(image_path):
        body_shape.width = Inches(4.5)
        slide.shapes.add_picture(image_path, Inches(5), Inches(2), width=Inches(4.5))

    return slide

# Portada
slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(slide_layout)
background = slide.background
fill = background.fill
fill.solid()
fill.fore_color.rgb = BG_COLOR
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Sistema Inteligente de Inventario"
title.text_frame.paragraphs[0].font.color.rgb = ACCENT_COLOR
title.text_frame.paragraphs[0].font.bold = True
title.text_frame.paragraphs[0].font.name = "Arial"
subtitle.text = "Propuesta Comercial para Coctelería Lovo\nPor: MDev - Soluciones Tecnológicas"
subtitle.text_frame.paragraphs[0].font.color.rgb = TEXT_COLOR
subtitle.text_frame.paragraphs[0].font.name = "Arial"

# Imágenes limpias
img1 = r"C:\Users\Snake\.gemini\antigravity-ide\brain\364b75f8-20d9-47ec-a117-26249bc64c86\clean_inventory_app_1789558727548.jpg"
img2 = r"C:\Users\Snake\.gemini\antigravity-ide\brain\364b75f8-20d9-47ec-a117-26249bc64c86\elegant_bartender_tablet_1789558740343.jpg"
img3 = r"C:\Users\Snake\.gemini\antigravity-ide\brain\364b75f8-20d9-47ec-a117-26249bc64c86\minimalist_chart_1789558753533.jpg"

# Diapositiva 2
add_slide(prs, "1. El Desafío vs La Solución", 
          "EL DESAFÍO:\n- Conteo manual (2-3 horas).\n- Mermas no detectadas.\n- Alto coste en horas.\n\nLA SOLUCIÓN:\n- Reconocimiento de Voz rápido.\n- Autocorrector Inteligente.\n- Agrupación por categorías.",
          img2)

# Diapositiva 3
add_slide(prs, "2. Funcionalidades Clave", 
          "- Interfaz POS Unificada: Registro fluido.\n- Gestión Multi-Ubicación: Barras y almacén.\n- Multiusuario: Trabajo en equipo en tiempo real.\n- Roles y Seguridad: Accesos controlados.\n- Diseño Limpio: Orientado a la usabilidad.",
          img1)

# Diapositiva 4
add_slide(prs, "3. Analítica y Control", 
          "- Historial Visual: Comparativas de consumo.\n- Exportación a Excel: Reportes claros.\n- Panel Autogestionado: Total independencia.\n- Cierre de Ciclo: Reseteo eficiente de base de datos.",
          img3)

# Diapositiva 5
add_slide(prs, "4. Modelos de Adquisición", 
          "MODELO A: Licencia Perpetua (5.500 €)\n- Pago único.\n\nMODELO B: SaaS / Suscripción (150 €/mes + Setup)\n- Bajo riesgo. Mantenimiento incluido.\n\nMODELO C: Exclusividad Absoluta (12.500 €)\n- Código fuente completo.")

# Diapositiva 6
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = BG_COLOR
slide.shapes.title.text = "¿Hablamos?"
slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = ACCENT_COLOR
slide.placeholders[1].text = "Quedo a su entera disposición para una reunión y demostración.\n\nMDev - Soluciones Tecnológicas"
slide.placeholders[1].text_frame.paragraphs[0].font.color.rgb = TEXT_COLOR

prs.save("Presentacion_Lovo.pptx")
print("Presentation generated successfully with light theme!")
